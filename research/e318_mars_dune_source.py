"""E318: read-only USGS Mars dune-source experiment, not E001 recovery.

Official USGS MGD3 merged release. Preserve raw source codes, provenance,
missingness and negative results. No inferred negative sites.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
from collections import Counter
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import Request, urlopen
from zipfile import ZipFile, BadZipFile

ITEM_ID = "5f8a1d2e82ce324187893a97"
METADATA_URL = f"https://www.sciencebase.gov/catalog/item/{ITEM_ID}?format=json"
MAX_DOWNLOAD = 32_000_000
MAX_EXPANDED = 12_000_000
MAX_ROWS = 20_000
MAX_TABLES = 24
EXTENSIONS = (".csv", ".dbf", ".xlsx", ".zip")


def allowed_file_url(url: str) -> bool:
    p = urlsplit(url)
    return (p.scheme == "https" and p.hostname in ("www.sciencebase.gov", "sciencebase.gov")
            and p.path.startswith("/catalog/file/get/") and not p.username and not p.password)


def read_public(url: str, limit: int) -> bytes:
    with urlopen(Request(url, headers={"User-Agent": "NASA-MGD3-E318/1.0"}), timeout=40) as response:
        target = urlsplit(response.geturl())
        if target.scheme != "https" or target.hostname not in ("www.sciencebase.gov", "sciencebase.gov"):
            raise ValueError("redirect outside official ScienceBase host")
        content = response.read(limit + 1)
    if len(content) > limit:
        raise ValueError("public download exceeds bound")
    return content


def select_files(item: dict) -> list[dict]:
    if item.get("id") != ITEM_ID or "dune" not in item.get("title", "").lower():
        raise ValueError("unexpected USGS Mars dune item identity")
    selected = []
    for f in item.get("files", []):
        name = str(f.get("name") or "")
        url = str(f.get("downloadUri") or "")
        if not name.lower().endswith(EXTENSIONS):
            continue
        size = f.get("size")
        selected.append({"name": name, "declared_bytes": size,
                         "uri": url if allowed_file_url(url) else None,
                         "eligible": bool(allowed_file_url(url) and isinstance(size, int)
                                          and 0 < size <= MAX_DOWNLOAD)})
    return selected


def _summary(rows, fields: list[str], name: str, format_name: str) -> dict:
    n = 0
    present = Counter()
    categories: dict[str, Counter] = {}
    for row in rows:
        n += 1
        if n > MAX_ROWS:
            raise ValueError("table row bound exceeded")
        for key, val in zip(fields, row):
            if val is None or not str(val).strip():
                continue
            present[key] += 1
            k = key.lower().replace("_", "").replace(" ", "")
            if ("dune" in k and "type" in k) or k in ("dunetype", "duneform"):
                if len(str(val)) <= 80:
                    categories.setdefault(key, Counter())[str(val).strip()] += 1
    return {"file": name, "format": format_name, "rows": n, "columns": fields[:120],
            "nonempty_by_field": dict(present),
            "dune_type_raw_codes": {k: dict(v.most_common(50)) for k, v in categories.items()},
            "codebook_interpretation": "UNVERIFIED; raw source codes are not harmonized labels"}


def profile_table(data: bytes, name: str) -> list[dict]:
    lower = name.lower()
    if len(data) > MAX_EXPANDED:
        raise ValueError("expanded member exceeds bound")
    if lower.endswith(".csv"):
        raw = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(raw))
        header = next(reader, [])
        if not header or len(header) > 120:
            raise ValueError("invalid CSV header")
        return [_summary(reader, header, name, "csv")]
    if lower.endswith(".dbf"):
        import shapefile
        table = shapefile.Reader(dbf=io.BytesIO(data), encoding="utf-8", encodingErrors="replace")
        header = [x[0] for x in table.fields[1:]]
        return [_summary(table.records(), header, name, "dbf")]
    if lower.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        try:
            for ws in wb.worksheets[:5]:
                records = ws.iter_rows(values_only=True)
                header = [str(x or "") for x in next(records, [])]
                if not header or len(header) > 120:
                    continue
                out.append(_summary(records, header, f"{name}:{ws.title}", "xlsx"))
        finally:
            wb.close()
        return out
    if lower.endswith(".zip"):
        out = []
        with ZipFile(io.BytesIO(data)) as z:
            members = [m for m in z.infolist() if not m.is_dir() and
                       m.filename.lower().endswith((".csv", ".dbf", ".xlsx"))]
            for m in members[:MAX_TABLES]:
                if m.file_size > MAX_EXPANDED:
                    out.append({"file": m.filename, "status": "OVERSIZED_TABLE"})
                    continue
                with z.open(m) as f:
                    raw = f.read(MAX_EXPANDED + 1)
                if len(raw) > MAX_EXPANDED:
                    raise ValueError("zip member decompression bound exceeded")
                try:
                    out.extend(profile_table(raw, m.filename))
                except (OSError, ValueError, BadZipFile, ImportError, UnicodeError) as exc:
                    out.append({"file": m.filename, "status": "TABLE_PARSE_FAILED",
                                "error_type": type(exc).__name__})
        return out
    return []


def audit(item: dict, downloader=read_public) -> dict:
    candidates = select_files(item)
    receipts = []
    for candidate in candidates:
        public = {k: v for k, v in candidate.items() if k != "uri"}
        if not candidate["eligible"]:
            public["status"] = "NOT_ELIGIBLE"
            receipts.append(public)
            continue
        try:
            raw = downloader(candidate["uri"], MAX_DOWNLOAD)
            public.update({"status": "DOWNLOADED", "sha256": hashlib.sha256(raw).hexdigest(),
                           "bytes": len(raw), "tables": profile_table(raw, candidate["name"])})
        except (OSError, ValueError, BadZipFile, ImportError, UnicodeError) as exc:
            public.update({"status": "ACCESS_OR_PARSE_FAILED", "error_type": type(exc).__name__})
        receipts.append(public)
    tables = [t for r in receipts for t in r.get("tables", []) if "rows" in t]
    return {"source": METADATA_URL, "item_id": ITEM_ID, "title": item.get("title"),
            "file_receipts": receipts, "tables_profiled": len(tables),
            "total_table_rows_not_unique_sites": sum(t["rows"] for t in tables),
            "status": "REAL_MARS_TABLES_PROFILED" if tables else "NO_MARS_TABLES_PROFILED",
            "limitations": ["MGD3 primarily inventories known positive dune fields; unmapped areas are not verified negatives.",
                            "Table rows may overlap across sheets and must not be counted as distinct sites.",
                            "Raw dune-type codes and PAAD Earth labels are not harmonized.",
                            "Coordinate frames, units, acquisition bias and sample support not yet reconciled."],
            "e001_known_site_recovery": "NOT YET MEASURED",
            "cross_body_similarity": "NOT YET PROVEN"}


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--output", type=Path, required=True)
    args = p.parse_args()
    raw = read_public(METADATA_URL, 2_000_000)
    result = audit(json.loads(raw))
    result["metadata_sha256"] = hashlib.sha256(raw).hexdigest()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({"status": result["status"], "tables_profiled": result["tables_profiled"],
                      "total_table_rows_not_unique_sites": result["total_table_rows_not_unique_sites"],
                      "files": [{"name": f["name"], "status": f["status"]} for f in result["file_receipts"]]}, sort_keys=True))


if __name__ == "__main__":
    main()
