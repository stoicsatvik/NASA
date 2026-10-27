"""E317: independent USGS PAAD aeolian analogue site-source access audit.

Read-only USGS ScienceBase item, bounded downloads, no auth, no site fabrication.
This is NOT E001 cross-body analogue recovery: PAAD describes Earth dune sites
but does not automatically provide matched Mars/Moon labels or negative controls.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import Request, urlopen
from zipfile import ZipFile, BadZipFile

ITEM_ID = "64d418c3d34ef477cf3ced10"
DOI = "10.5066/P9IFVZEX"
SOURCE = f"https://www.sciencebase.gov/catalog/item/{ITEM_ID}?format=json"
MAX_BYTES = 25_000_000
MAX_MEMBER = 10_000_000
MAX_ROWS = 5000


def _safe_url(url: str) -> bool:
    p = urlsplit(url)
    return p.scheme == "https" and p.hostname in {
        "www.sciencebase.gov", "sciencebase.gov"
    } and p.path.startswith("/catalog/file/get/") and not p.username and not p.password


def _read(url: str, limit: int = MAX_BYTES) -> bytes:
    with urlopen(Request(url, headers={"User-Agent": "NASA-PAAD-E317/1.0"}), timeout=35) as response:
        if urlsplit(response.geturl()).hostname not in ("www.sciencebase.gov", "sciencebase.gov"):
            raise ValueError("unexpected download redirect")
        data = response.read(limit + 1)
    if len(data) > limit:
        raise ValueError("bounded public download exceeded")
    return data


def metadata_summary(item: dict) -> dict:
    if item.get("id") != ITEM_ID or "Planetary Aeolian Analog Database" not in item.get("title", ""):
        raise ValueError("USGS item identity mismatch")
    files = []
    for f in item.get("files", []):
        uri = f.get("downloadUri") or ""
        files.append({
            "name": f.get("name"), "bytes_declared": f.get("size"),
            "allowed_public_uri": _safe_url(uri),
            "_uri": uri if _safe_url(uri) else None,
        })
    return {"item_id": ITEM_ID, "doi": DOI, "title": item["title"], "files": files}


def _csv_profile(data: bytes, name: str) -> dict:
    text = data.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = rows.fieldnames or []
    if not headers or len(headers) > 100:
        raise ValueError("invalid table headers")
    count = 0
    for row in rows:
        if not any(v and v.strip() for v in row.values() if isinstance(v, str)):
            continue
        count += 1
        if count > MAX_ROWS:
            raise ValueError("too many rows")
    return {"file": name, "format": "csv", "rows": count, "columns": headers}


def _xlsx_profile(data: bytes, name: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets[:8]:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, ())
        columns = [str(v) if v is not None else "" for v in header]
        count = 0
        for row in rows:
            if any(v is not None and str(v).strip() for v in row):
                count += 1
            if count > MAX_ROWS:
                raise ValueError("too many rows")
        sheets.append({"sheet": ws.title, "rows": count, "columns": columns[:100]})
    wb.close()
    return {"file": name, "format": "xlsx", "sheets": sheets}


def _dbf_profile(data: bytes, name: str) -> dict:
    import shapefile
    reader = shapefile.Reader(dbf=io.BytesIO(data), encoding="utf-8", encodingErrors="replace")
    return {"file": name, "format": "dbf", "rows": len(reader),
            "columns": [f[0] for f in reader.fields[1:]]}


def profile(data: bytes, name: str) -> list[dict]:
    low = name.lower()
    if low.endswith(".csv"):
        return [_csv_profile(data, name)]
    if low.endswith(".xlsx"):
        return [_xlsx_profile(data, name)]
    if low.endswith(".dbf"):
        return [_dbf_profile(data, name)]
    if not low.endswith(".zip"):
        return []
    out = []
    with ZipFile(io.BytesIO(data)) as archive:
        members = [i for i in archive.infolist()
                   if not i.is_dir() and i.filename.lower().endswith((".csv", ".xlsx", ".dbf"))]
        for member in members[:30]:
            if member.file_size > MAX_MEMBER:
                out.append({"file": member.filename, "status": "OVERSIZED_MEMBER"})
                continue
            content = archive.read(member)
            if len(content) > MAX_MEMBER:
                raise ValueError("expanded member exceeds bound")
            out.extend(profile(content, member.filename))
    return out


def audit(item: dict, downloader=_read) -> dict:
    summary = metadata_summary(item)
    profiles = []
    for f in summary["files"]:
        uri = f.pop("_uri")
        name = f.get("name") or ""
        if not uri or not name.lower().endswith((".zip", ".csv", ".xlsx", ".dbf")):
            continue
        if f.get("bytes_declared") and f["bytes_declared"] > MAX_BYTES:
            profiles.append({"file": name, "status": "DECLARED_TOO_LARGE"})
            continue
        try:
            raw = downloader(uri)
            profiles.append({"file": name, "sha256": hashlib.sha256(raw).hexdigest(),
                             "bytes": len(raw), "tables": profile(raw, name)})
        except (OSError, ValueError, BadZipFile, ImportError, UnicodeError) as exc:
            profiles.append({"file": name, "status": "ACCESS_OR_PARSE_FAILED",
                             "error_type": type(exc).__name__})
    summary["file_profiles"] = profiles
    summary["site_rows_observed"] = sum(
        table.get("rows", 0) for p in profiles for table in p.get("tables", [])
        if table.get("format") in ("csv", "dbf")
    )
    summary["status"] = "TABULAR_SOURCE_PROFILED" if any(p.get("tables") for p in profiles) else "NO_TABULAR_SOURCE_PROFILED"
    summary["claim"] = "Source access and table schema only; Earth dune analog labels NOT cross-body matches"
    summary["e001_known_site_recovery"] = "NOT YET MEASURED"
    summary["cross_body_similarity"] = "NOT YET PROVEN"
    summary["limitations"] = [
        "Published count of 45 Earth dune sites not independently equated to parsed row count.",
        "No blind positive/negative Moon/Mars matches or feature comparability established.",
        "Metadata/table headers do not prove scientific accuracy of records.",
        "No private credentials, raw datasets, field measurements or submissions.",
    ]
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    raw = _read(SOURCE, 2_000_000)
    result = audit(json.loads(raw))
    result["metadata_sha256"] = hashlib.sha256(raw).hexdigest()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n")
    print(json.dumps({"status": result["status"], "item_id": ITEM_ID,
                      "site_rows_observed": result["site_rows_observed"],
                      "file_profiles": result["file_profiles"]}, sort_keys=True))


if __name__ == "__main__":
    main()
